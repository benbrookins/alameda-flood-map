"""Spreadsheet of potential shelters to contact first, for Bay +2 ft and +3 ft.

Uses the shelter assignment from 11_capacity_gaps.py (2 km walk along open paths, flooded paths closed). Potential sites are listed when they are among
the suggested contacts for (a) a pre-identified shelter that is over or near capacity (75%+ at the default shelter-use
share) or (b) an area with no pre-identified shelter in range. People counts that depend on the shelter-use share are
Excel formulas tied to one input cell, so the share can be changed in the workbook.

Output (contains capacities, so kept out of the public repo): data/work/outreach_priority_2ft_3ft.xlsx

With --sites faith-community the candidate list is searched afresh for places of worship, community, recreation and
senior centers, and libraries only (no schools), up to 8 per problem area, ranked by reach, building size, then walking
distance. Output: data/work/outreach_faith_community_2ft_3ft.xlsx
"""
import argparse
import json
from collections import defaultdict

import geopandas as gpd
import numpy as np

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import DOCS, WORK, load_access, scenario_mask, shelter_sites

SCENARIOS = {"+2 ft": "b2_r0_c0", "+3 ft": "b3_r0_c0"}
KM = "2"
SHARE = 0.20
NEAR = 0.75
OUT = WORK / "outreach_priority_2ft_3ft.xlsx"
OUT_FAITH = WORK / "outreach_faith_community_2ft_3ft.xlsx"
FAITH_BUCKETS = {"worship", "community_center", "senior_center", "library"}
FAITH_KINDS = {"worship", "community", "library"}
UNNAMED = {"School", "Community center", "Library", "Place of worship"}
SAME_SITE_M = 250  # as in 11_capacity_gaps.py
PER_AREA = 8
BUCKET = {
    "college": "College or university", "high_school": "High school", "middle_school": "Middle school",
    "elementary_school": "Elementary school", "small_school": "Small or alternative school",
    "community_center": "Community or recreation center", "senior_center": "Senior center", "library": "Library",
    "worship": "Place of worship",
}
REL = {"L": "Larger than typical", "T": "Typical", "S": "Smaller than typical"}
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
INPUT_FONT, INPUT_FILL = Font(name="Arial", size=10, color="0000FF", bold=True), PatternFill("solid", fgColor="FFFF00")
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
SHARE_CELL = "'Read me'!$B$6"


def floods_at(mask):
    for ft in range(1, 5):
        if mask & (1 << (ft - 1)):
            return f"Bay +{ft} ft"
    return "Not at +1 to +4 ft"


def table(ws, headers, rows, widths):
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def faith_candidates(feats, gaps, cap, areas):
    """Fresh search: for each problem area, rebuild its flooded blocks per scenario (as 11_capacity_gaps.py assigns them),
    then rank worship / community / library sites within a 2 km walk of those blocks. Fills area["cands"] and area["reach"]."""
    z = np.load(WORK / "block_flood.npz")
    keys, F = list(z["keys"]), z["pop"].astype(np.float64)
    tr = gpd.read_file(DOCS / "tracts.geojson")
    tract_city = dict(zip(tr.GEOID.str[5:], tr.place))
    city = np.array([tract_city.get(t, "Unincorporated") for t in z["tract"]])
    props = {int(i): f["properties"] for i, f in feats.items()}
    ids, sx, sy = shelter_sites()
    row_of = {int(i): r for r, i in enumerate(ids)}
    pre = [i for i in props if props[i]["k"] == "pre" and str(i) in cap["sites"]]
    pre_row = np.array([row_of[i] for i in pre])
    pre_cap = np.array([cap["sites"][str(i)][0] for i in pre])
    pre_xy = np.c_[sx[pre_row], sy[pre_row]]
    pot = [i for i in props if props[i]["k"] in FAITH_KINDS and props[i].get("b") in FAITH_BUCKETS and props[i]["n"] not in UNNAMED
           and i in row_of]
    pot_row = np.array([row_of[i] for i in pot])
    all_pre_xy = np.c_[sx[[row_of[i] for i in props if props[i]["k"] == "pre" and i in row_of]],
                       sy[[row_of[i] for i in props if props[i]["k"] == "pre" and i in row_of]]]
    near_pre = np.hypot(sx[pot_row][:, None] - all_pre_xy[:, 0][None], sy[pot_row][:, None] - all_pre_xy[:, 1][None]).min(axis=1) < SAME_SITE_M
    size = json.load(open(DOCS / "site_size.json"))
    size_rank = np.array([{"L": 0, "T": 1, "?": 1, "S": 2}.get(size.get(str(i), [0, "?"])[1], 1) for i in pot])
    masks = np.array([props[i]["m"] for i in pot])
    R = float(KM) * 1000

    reach = defaultdict(lambda: defaultdict(lambda: np.zeros(len(pot))))   # city -> label -> reach per site
    mind = defaultdict(lambda: np.full(len(pot), np.inf))                  # city -> nearest walk to any area block
    for label, key in SCENARIOS.items():
        fl = F[keys.index(key)]
        t = load_access(key)
        D = np.full((len(ids), len(fl)), np.inf, np.float32)
        D[t["walk_site"], t["walk_block"]] = t["walk_m"]
        dry_pre = np.array([(props[i]["m"] & scenario_mask(key)) == 0 for i in pre])
        Dd = np.where(dry_pre[None, :], D[pre_row].T, np.inf)
        nearest, dist = Dd.argmin(axis=1), Dd.min(axis=1)
        flooded = fl > 0
        covered = flooded & (dist <= R)
        dry_pot = (masks & scenario_mask(key)) == 0
        # the assignment must match what the dashboard and the status tab use
        check = np.bincount(nearest[covered], weights=fl[covered], minlength=len(pre))
        for j, i in enumerate(pre):
            assert round(check[j]) == gaps[key][KM]["a"].get(str(i), 0), (key, i)
        blocks = {}
        for j, i in enumerate(pre):
            if check[j] * SHARE / pre_cap[j] >= NEAR:
                blocks.setdefault(props[i]["c"], np.zeros(len(fl), bool))
                blocks[props[i]["c"]] |= covered & (nearest == j)
        for cty, (n, _, _) in gaps[key][KM]["u"].items():
            blocks.setdefault(cty, np.zeros(len(fl), bool))
            blocks[cty] |= flooded & ~covered & (city == cty)
        for cty, sel in blocks.items():
            if not sel.any():
                continue
            near = (D[pot_row][:, sel] <= R)
            r = near @ fl[sel]
            r[~dry_pot | near_pre] = 0
            reach[cty][label] = r
            d = np.where(dry_pot & ~near_pre, D[pot_row][:, sel].min(axis=1), np.inf)
            mind[cty] = np.minimum(mind[cty], d)
    for cty, area in areas.items():
        tot = sum(reach[cty][l] for l in SCENARIOS) if cty in reach else np.zeros(len(pot))
        order = sorted((k for k in range(len(pot)) if tot[k] >= 1),
                       key=lambda k: (-round(tot[k]), size_rank[k], mind[cty][k]))[:PER_AREA]
        area["cands"] = {pot[k]: n for n, k in enumerate(order)}
        area["reach"] = defaultdict(lambda: defaultdict(int))
        for k in order:
            for l in SCENARIOS:
                area["reach"][pot[k]][l] = int(round(reach[cty][l][k])) if l in reach[cty] else 0


def main(faith=False):
    gaps = json.load(open(DOCS / "capacity_gaps.json"))
    cap = json.load(open(DOCS / "capacity.json"))
    size = json.load(open(DOCS / "site_size.json"))
    feats = {f["properties"]["id"]: f for f in json.load(open(DOCS / "points.json"))["features"]}
    name = lambda i: feats[int(i)]["properties"]["n"]

    # Problem areas by city: shortfalls at over/near-capacity pre-identified shelters, and residents with none in range.
    areas = defaultdict(lambda: {"need": defaultdict(float), "problems": defaultdict(list), "cands": {}, "reach": defaultdict(lambda: defaultdict(int))})
    status_rows, uncovered_rows = {}, []
    for label, key in SCENARIOS.items():
        e = gaps[key][KM]
        for sid, assigned in e["a"].items():
            c = cap["sites"][sid][0]
            status_rows.setdefault(sid, {})[label] = assigned
            load = assigned * SHARE / c
            if load >= NEAR:
                city = feats[int(sid)]["properties"]["c"]
                area = areas[city]
                area["need"][label] += max(0.0, assigned * SHARE - c)
                area["problems"][label].append(f"{name(sid)} {'over' if load >= 1 else 'near'} capacity ({load:.0%})")
                for cid, n in e["c"].get(sid, []):
                    area["cands"].setdefault(cid, len(area["cands"]))
                    area["reach"][cid][label] += n
        for city, (n, cands, _) in e["u"].items():
            uncovered_rows.append([label, city, n, None, "; ".join(name(c) for c, _ in cands[:5])])
            area = areas[city]
            area["need"][label] += n * SHARE
            area["problems"][label].append(f"~{round(n * SHARE)} people with no pre-identified shelter within a {KM} km walk")
            for cid, r in cands:
                area["cands"].setdefault(cid, len(area["cands"]))
                area["reach"][cid][label] += r
    if faith:
        faith_candidates({int(i): f for i, f in feats.items()}, gaps, cap, areas)
    ranked_areas = sorted(areas.items(), key=lambda kv: (-kv[1]["need"]["+3 ft"], -kv[1]["need"]["+2 ft"]))

    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    notes = [
        ["Priority potential shelters to contact: Bay +2 ft and +3 ft", None],
        ["Prepared 2026-10-02 from the Alameda County Flood Risk Explorer data (estimates).", None],
        [None, None],
        ["Settings", None],
        ["Flood scenarios", "Bay water level +2 ft and +3 ft above normal high tide (NOAA maps); FEMA zones and low-lying areas off"],
        ["Share of affected residents needing a public shelter", SHARE],
        ["Distance to a shelter", f"{KM} km walk (about 30 minutes) along open paths; flooded paths are closed"],
        [None, None],
        ["How to read this", None],
        ["Priority contacts", "Grouped by problem area (city), most severe first: pre-identified shelters over or near capacity, and flooded residents with no pre-identified shelter in range. Within each area, potential sites (not yet confirmed as shelters) are ordered by affected residents in reach at +2 and +3 ft combined; sites that themselves flood by +3 ft are listed last."],
        ["Pre-identified status", "How many people each pre-identified shelter would receive (nearest dry shelter within range) versus its overnight capacity."],
        ["Uncovered areas", f"Flooded residents with no dry pre-identified shelter within a {KM} km walk, by city."],
        ["Edit", "Change the yellow cell (B6) to try a different shelter-use share; people counts and statuses update. The list of sites was chosen at 20%."],
        [None, None],
        ["Caveats", None],
        ["Estimated capacity", "Typical overnight capacity of pre-identified shelters of the same type; can be off by a third or more. No estimate for elementary schools, libraries, small schools."],
        ["Building size", "Main building footprint from map outlines, compared with other sites of the same type."],
        ["Reach", "Flooded residents within the distance of the site who would otherwise go to the over-capacity shelter or have no shelter in range. Sites near each other reach many of the same people, so reach is not additive across sites."],
        ["Status", "Sites have not been contacted. Not an official plan."],
    ]
    if faith:
        notes[0][0] = "Faith-based, community center and library outreach: Bay +2 ft and +3 ft (no schools)"
        notes[9][1] = ("Grouped by problem area (city), most severe first: pre-identified shelters over or near capacity, and flooded residents "
                       "with no pre-identified shelter in range. Within each area, up to 8 places of worship, community, recreation and senior "
                       "centers, and libraries (not yet confirmed as shelters) that are dry at the level shown, ordered by affected residents "
                       "in reach at +2 and +3 ft combined, then building size, then walking distance.")
        notes.append(["Rough capacity estimates", "Estimated capacities for these site types are rough, based on 7-23 surveyed sites per type and "
                      "none for libraries, so the building size column matters more than usual."])
    for r in notes:
        readme.append(r)
    for row in readme.iter_rows():
        for c in row:
            c.font = FONT
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in (1, 4, 9, 15):
        readme.cell(r, 1).font = Font(name="Arial", size=12 if r == 1 else 10, bold=True)
    share = readme["B6"]
    share.font, share.fill, share.number_format = INPUT_FONT, INPUT_FILL, "0%"
    share.comment = Comment("Planning assumption (placeholder): the share of residents in flooded areas who would need a "
                            "public shelter. 20% chosen by the user pending guidance.", "dashboard")
    readme.column_dimensions["A"].width, readme.column_dimensions["B"].width = 34, 110

    # Priority contacts, grouped by problem area (most severe first); sites in the dashboard's suggested order
    ws = wb.create_sheet("Priority contacts")
    rows = []
    for area_rank, (city, area) in enumerate(ranked_areas, 1):
        problems = ["; ".join(area["problems"]["+2 ft"]) or "No issue at +2 ft", "; ".join(area["problems"]["+3 ft"]) or "No issue at +3 ft"]
        if not area["cands"]:
            rows.append([area_rank, city, *problems, "-", f"No dry potential site within a {KM} km walk of these residents",
                         "", "", "", 0, 0, "", "", "", ""])
            continue
        floods_by_3ft = lambda c: bool(feats[c]["properties"]["m"] & 0b0111)  # Bay +1, +2 or +3 ft
        order = sorted(area["cands"], key=lambda c: (False if faith else floods_by_3ft(c), -(area["reach"][c]["+2 ft"] + area["reach"][c]["+3 ft"]), area["cands"][c]))
        for site_rank, cid in enumerate(order, 1):
            p = feats[cid]["properties"]
            lon, lat = feats[cid]["geometry"]["coordinates"]
            t = cap["types"].get(p.get("b"))
            rel = size.get(str(cid), [0, "?"])[1]
            r = len(rows) + 2
            rows.append([area_rank, city, *problems, site_rank, p["n"],
                         BUCKET.get(p.get("b"), p["k"]), t["overnight"] if t else "No estimate", REL.get(rel, "Unknown"),
                         area["reach"][cid]["+2 ft"], area["reach"][cid]["+3 ft"],
                         f"=ROUND(J{r}*{SHARE_CELL},0)", f"=ROUND(K{r}*{SHARE_CELL},0)",
                         floods_at(p["m"]), f'=HYPERLINK("https://www.google.com/maps?q={lat},{lon}","Map")'])
    table(ws, ["Area priority", "Area (city)", "Problem at +2 ft", "Problem at +3 ft", "Site order in area", "Site to contact",
               "Type", "Est. overnight capacity", "Building size (vs. same type)",
               "Affected residents in reach, +2 ft (all)", "Affected residents in reach, +3 ft (all)",
               "Likely to need shelter, +2 ft", "Likely to need shelter, +3 ft", "Site itself floods at", "Map"],
          rows, [8, 13, 38, 38, 8, 34, 20, 11, 14, 12, 12, 11, 11, 14, 7])

    # Pre-identified status
    ws = wb.create_sheet("Pre-identified status")
    srt = sorted(status_rows, key=lambda s: -max(v * SHARE / cap["sites"][s][0] for v in status_rows[s].values()))
    rows = []
    for i, sid in enumerate(srt, 2):
        p = feats[int(sid)]["properties"]
        rows.append([p["n"], p["c"], cap["sites"][sid][0], status_rows[sid].get("+2 ft", 0), status_rows[sid].get("+3 ft", 0),
                     f"=ROUND(D{i}*{SHARE_CELL},0)", f"=ROUND(E{i}*{SHARE_CELL},0)",
                     f'=IF(C{i}>0,F{i}/C{i},"")', f'=IF(C{i}>0,G{i}/C{i},"")',
                     f'=IF(H{i}="","",IF(H{i}>=1,"Over capacity",IF(H{i}>={NEAR},"Near capacity","Has room")))',
                     f'=IF(I{i}="","",IF(I{i}>=1,"Over capacity",IF(I{i}>={NEAR},"Near capacity","Has room")))'])
    table(ws, ["Pre-identified shelter", "City", "Overnight capacity", "Affected residents assigned, +2 ft (all)",
               "Affected residents assigned, +3 ft (all)", "Likely to need shelter, +2 ft", "Likely to need shelter, +3 ft",
               "Share of capacity used, +2 ft", "Share of capacity used, +3 ft", "Status, +2 ft", "Status, +3 ft"],
          rows, [36, 13, 11, 14, 14, 12, 12, 12, 12, 14, 14])
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=9):
        for c in row:
            c.number_format = "0%"

    # Uncovered areas
    ws = wb.create_sheet("Uncovered areas")
    uncovered_rows.sort(key=lambda r: (r[0], -r[2]))
    for i, r in enumerate(uncovered_rows, 2):
        r[3] = f"=ROUND(C{i}*{SHARE_CELL},0)"
    table(ws, ["Scenario", "City", f"Flooded residents with no pre-identified shelter within a {KM} km walk (all)",
               "Likely to need shelter", "Suggested sites to contact"], uncovered_rows, [10, 14, 20, 12, 90])

    rows_written = list(wb["Priority contacts"].iter_rows(min_row=2))
    out = OUT_FAITH if faith else OUT
    wb.save(out)
    print(f"wrote {out} | {len(ranked_areas)} problem areas, {len(rows_written)} site rows, {len(srt)} pre-identified shelters, {len(uncovered_rows)} uncovered areas")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--sites", choices=["faith-community"], help="search only worship, community and library sites (no schools)")
    main(faith=ap.parse_args().sites == "faith-community")
