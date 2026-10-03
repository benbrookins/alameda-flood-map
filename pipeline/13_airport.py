"""Airport-area deep dive: flooding in 0.1 ft steps from +2.0 to +3.0 ft, spill points, ART cross-check, road cut-offs.

Grid: NOAA 3 m SF Bay lidar DEM (m NAVD88) and NOAA's MHHW tidal surface, read for the focus box only.
Water reach: for every cell, the lowest Bay level (ft above local MHHW) at which water from the open Bay can reach it
(a "bathtub" fill computed by morphological reconstruction). Flooded at level L means reach <= L; water depth is
L - ground (both in ft above local MHHW).

Stages (cached in data/work/airport/; pass a stage name to rerun from it): grid, reach, barriers, art, roads, publish.
"""
import json
import sys
import time

import geopandas as gpd
import numpy as np
import rasterio
import shapely
import rasterio.features
from rasterio.transform import Affine
from rasterio.warp import Resampling, reproject
from rasterio.windows import from_bounds
from scipy import ndimage
from skimage.morphology import reconstruction

from scipy.sparse import coo_matrix
from scipy.sparse.csgraph import connected_components
from scipy.spatial import cKDTree

from common import CRS, DOCS, FLOOD, WORK, block_shapes, road_network

OUT = WORK / "airport"
WEB = DOCS.parent / "airport" / "data"
FOCUS = (-122.345, 37.69, -122.15, 37.80)  # west, south, east, north
LEVELS = [round(2.0 + i / 10, 1) for i in range(11)]
NOAA = "/vsicurl/https://coast.noaa.gov/slrdata"
DEM_URL = f"{NOAA}/DEMs/CA/CA_SFBay_GCS_3m_NAVDm.tif"
MHHW_URL = f"{NOAA}/Tidal_Surfaces/WC_MHHW_GCS_100m_NAVDm.tif"
FT = 0.3048
ART_SERVICE = "https://geodata.dot.ca.gov/arcgis/rest/services/CHhqenvi/DEA_BCDC_polygon_SLR/FeatureServer"
ART_LAYERS = {12: 0, 24: 1, 36: 2, 52: 4}  # inches above MHHW -> layer id (48 in is empty in this service)
LAND_FT = 0.15     # ground at least this far above MHHW counts as land
BIG_SPILL_KM2 = 0.5
MAX_WALLS = 15
DEPTH_LIMIT_FT = 0.5   # a road is impassable where water on it is deeper than this
SAMPLE_M = 10          # spacing of depth samples along each road segment
SNAP_M = 400           # blocks and sites attach to the nearest road node within this distance
MIN_AREA_RESIDENTS = 1
POTENTIAL = {"school", "community", "library", "worship"}
SUBAREAS = {  # names for cut-off areas, by where most of their residents are; listed specific-first
    "Bay Farm Island": (-122.262, 37.722, -122.224, 37.748),
    "Alameda Point / Mariner Square": (-122.33, 37.778, -122.27, 37.80),
    "Oakland airport": (-122.235, 37.69, -122.195, 37.735),
    "Alameda main island": (-122.345, 37.748, -122.225, 37.80),
    "Coliseum / Hegenberger": (-122.225, 37.72, -122.17, 37.765),
    "San Leandro Marina": (-122.20, 37.69, -122.15, 37.72),
}
VARIANTS = {"mapped": "reach.npy", "hold": "reach_hold.npy"}
WATER_FT = -99.0  # relative elevation assigned to open water and NOAA's water mask


def log(t0, msg):
    print(f"{time.time() - t0:7.1f}s  {msg}", flush=True)


# ---- grid ------------------------------------------------------------------------------------

def build_grid():
    with rasterio.open(DEM_URL) as ds:
        win = from_bounds(*FOCUS, ds.transform)
        dem = ds.read(1, window=win).astype("float32")
        transform, crs, nodata = ds.window_transform(win), ds.crs, ds.nodata
    mhhw = np.zeros_like(dem)
    with rasterio.open(MHHW_URL) as ds:
        reproject(rasterio.band(ds, 1), mhhw, dst_transform=transform, dst_crs=crs, resampling=Resampling.bilinear)
    # Elevation relative to local MHHW, in feet. NOAA marks water as -99 m; nodata (rare) is treated as high ground.
    rel = (dem - mhhw) / FT
    rel[dem <= -98] = WATER_FT
    rel[dem == nodata] = 99.0
    np.savez_compressed(OUT / "grid.npz", rel=rel.astype("float32"), transform=np.array(transform)[:6], crs=str(crs))
    return rel, transform


def load_grid():
    z = np.load(OUT / "grid.npz")
    return z["rel"], Affine(*z["transform"])


def cell_km2(transform):
    lat = (FOCUS[1] + FOCUS[3]) / 2
    return abs(transform.a) * 111320 * np.cos(np.radians(lat)) * abs(transform.e) * 110950 / 1e6


# ---- water reach -----------------------------------------------------------------------------

def open_water(rel):
    """Water cells connected to the edge of the box (the Bay and channels), not inland ponds."""
    water = rel <= WATER_FT
    lab, _ = ndimage.label(water)
    edge = np.unique(np.r_[lab[0], lab[-1], lab[:, 0], lab[:, -1]])
    return np.isin(lab, edge[edge > 0])


def water_reach(rel, bay, eight=True):
    """Lowest Bay level (ft above MHHW) at which water from `bay` reaches each cell."""
    z = np.where(rel <= WATER_FT, -5.0, rel)  # inland ponds: low ground that fills once reached
    z = np.minimum(z, 50.0)
    seed = np.where(bay, z, z.max())
    footprint = np.ones((3, 3)) if eight else ndimage.generate_binary_structure(2, 1)
    return reconstruction(seed, z, method="erosion", footprint=footprint).astype("float32")


def calibrate(rel, transform, bay):
    """Compare computed flooding with NOAA's published polygons on land, for 4- and 8-neighbor connectivity."""
    land = rel > 0.15
    k = cell_km2(transform)
    results = {}
    for eight in (True, False):
        reach = water_reach(rel, bay, eight)
        rows = {}
        for ft in (2.0, 2.5, 3.0):
            poly = gpd.read_file(FLOOD, layer=f"bay_{ft:g}ft").to_crs(4326).geometry.iloc[0]
            noaa = rasterio.features.rasterize([(poly, 1)], out_shape=rel.shape, transform=transform, fill=0, dtype="uint8").astype(bool)
            ours = reach <= ft
            a, b = (ours & land), (noaa & land)
            rows[ft] = {"ours_km2": round(a.sum() * k, 2), "noaa_km2": round(b.sum() * k, 2),
                        "iou": round((a & b).sum() / max((a | b).sum(), 1), 3)}
        results["8-neighbor" if eight else "4-neighbor"] = rows
    return results


# ---- barriers ------------------------------------------------------------------------------

def find_spills(reach, rel, transform, min_km2):
    """Spill points between +2.0 and +3.0 ft: where water first gets into each newly flooded area of land."""
    land, k, st = rel > LAND_FT, cell_km2(transform), np.ones((3, 3))
    out, prev = [], 1.9
    for L in LEVELS:
        before, new = reach <= prev, (reach > prev) & (reach <= L)
        lab, n = ndimage.label(new, structure=st)
        if n:
            area = ndimage.sum(new & land, lab, range(1, n + 1)) * k
            for c in np.argsort(-area) + 1:
                if area[c - 1] < min_km2:
                    break
                rr, cc = np.nonzero((lab == c) & ndimage.binary_dilation(before, structure=st))
                if len(rr):
                    j = np.argmax(reach[rr, cc])
                    lon, lat = transform * (cc[j] + 0.5, rr[j] + 0.5)
                    out.append({"level": L, "crest_ft": round(float(reach[rr[j], cc[j]]), 2), "land_km2": round(float(area[c - 1]), 2),
                                "lat": round(lat, 5), "lon": round(lon, 5), "row": int(rr[j]), "col": int(cc[j])})
        prev = L
    return out


def barriers_hold(rel, transform, bay):
    """Block spill points that each flood >= BIG_SPILL_KM2 of land, one at a time, until none remain below +3 ft.

    Blocking raises crest-height ground (1.0-3.5 ft above MHHW) within 150 m of the spill point above the range,
    as if that stretch of shoreline protection held."""
    rel2, walls = rel.copy(), []
    for _ in range(MAX_WALLS):
        spills = find_spills(water_reach(rel2, bay), rel2, transform, BIG_SPILL_KM2)
        if not spills:
            break
        s = min(spills, key=lambda x: x["crest_ft"])
        r, c, rad = s["row"], s["col"], 50
        sub = rel2[max(r - rad, 0):r + rad + 1, max(c - rad, 0):c + rad + 1]
        sub[(sub > 1.0) & (sub < 3.5)] = 5.0
        walls.append({k: v for k, v in s.items() if k not in ("row", "col")})
        print(f"    blocked spill at +{s['crest_ft']} ft ({s['land_km2']} km2) near {s['lat']}, {s['lon']}", flush=True)
    return water_reach(rel2, bay), walls


# ---- ART ---------------------------------------------------------------------------------------

def fetch_art():
    """ART flood polygons for the focus box, saved as Esri JSON (keeps holes); fetched one feature at a time."""
    import requests
    from common import HEADERS
    w, s_, e, n = FOCUS
    for inch, layer in ART_LAYERS.items():
        path = OUT / f"art_{inch}in.json"
        if path.exists():
            continue
        feats, offset = None, 0
        while True:
            r = requests.get(f"{ART_SERVICE}/{layer}/query", headers=HEADERS, timeout=600, params={
                "where": "1=1", "outFields": "DEPTH_FT", "f": "json", "geometry": f"{w},{s_},{e},{n}",
                "geometryType": "esriGeometryEnvelope", "inSR": 4326, "spatialRel": "esriSpatialRelIntersects",
                "outSR": 4326, "geometryPrecision": 6, "resultOffset": offset, "resultRecordCount": 1})
            r.raise_for_status()
            page = r.json()
            if feats is None:
                feats = page
            elif page.get("features"):
                feats["features"] += page["features"]
            if not page.get("features"):
                break
            offset += 1
            print(f"    ART {inch} in: {offset} features", flush=True)
        path.write_text(json.dumps(feats))


def art_grids(shape, transform):
    out = {}
    for inch in ART_LAYERS:
        g = gpd.read_file(OUT / f"art_{inch}in.json", driver="ESRIJSON").to_crs(4326)
        out[inch] = rasterio.features.rasterize(((geom, 1) for geom in g.geometry if geom is not None), out_shape=shape,
                                                transform=transform, fill=0, dtype="uint8").astype(bool)
    return out


def compare_art(rel, transform, reach, reach_hold):
    land, k = rel > LAND_FT, cell_km2(transform)
    grids = art_grids(rel.shape, transform)
    rows = []
    for inch, art in grids.items():
        ft = inch / 12
        for name, rc in (("as mapped", reach), ("barriers hold", reach_hold)):
            noaa = rc <= ft
            a, b = art & land, noaa & land
            rows.append({"inches": inch, "ft": round(ft, 2), "variant": name, "art_km2": round(a.sum() * k, 2),
                         "noaa_km2": round(b.sum() * k, 2), "both_km2": round((a & b).sum() * k, 2),
                         "iou": round((a & b).sum() / max((a | b).sum(), 1), 3)})
    np.savez_compressed(OUT / "art_grids.npz", **{f"in{inch}": g for inch, g in grids.items()})
    return rows


# ---- roads and cut-off areas -----------------------------------------------------------------

def subarea(lon, lat, weights=None):
    w = np.ones(len(lon)) if weights is None else np.asarray(weights, float)
    best, score = "Other", 0.0
    for name, (x0, y0, x1, y1) in SUBAREAS.items():
        inside = w[(lon >= x0) & (lon <= x1) & (lat >= y0) & (lat <= y1)].sum()
        if inside > score:
            best, score = name, inside
    return best


def edge_cut_levels(net, rel, transform, reach):
    """Bay level at which each road segment becomes impassable (inf if never within the focus box)."""
    a, b = net["a"], net["b"]
    w, s_, e, n = FOCUS
    lon, lat = net["lon"], net["lat"]
    inside = lambda i: (lon[i] >= w) & (lon[i] <= e) & (lat[i] >= s_) & (lat[i] <= n)
    idx = np.flatnonzero(net["at_grade"] & (inside(a) | inside(b)))
    length = np.hypot(net["x"][a[idx]] - net["x"][b[idx]], net["y"][a[idx]] - net["y"][b[idx]])
    k = np.maximum(2, np.ceil(length / SAMPLE_M).astype(int) + 1)
    owner = np.repeat(np.arange(len(idx)), k)
    t = np.concatenate([np.linspace(0, 1, m) for m in k])
    slon = lon[a[idx]][owner] * (1 - t) + lon[b[idx]][owner] * t
    slat = lat[a[idx]][owner] * (1 - t) + lat[b[idx]][owner] * t
    cols, rows = ~transform * (slon, slat)
    rows, cols = np.floor(rows).astype(int), np.floor(cols).astype(int)
    ok = (rows >= 0) & (rows < rel.shape[0]) & (cols >= 0) & (cols < rel.shape[1])
    cell = np.full(len(owner), np.inf, "float32")
    r_, c_ = rows[ok], cols[ok]
    level = np.maximum(reach[r_, c_], rel[r_, c_] + DEPTH_LIMIT_FT)
    level[rel[r_, c_] <= WATER_FT] = np.inf  # samples falling on mapped water are position error, not flooding
    cell[ok] = level
    per_edge = np.full(len(idx), np.inf, "float32")
    np.minimum.at(per_edge, owner, cell)
    out = np.full(len(a), np.inf, "float32")
    out[idx] = per_edge
    return out


def components(n, a, b, keep):
    g = coo_matrix((np.ones(int(keep.sum())), (a[keep], b[keep])), shape=(n, n))
    return connected_components(g, directed=False)[1]


def edges_outline(net, edge_idx):
    x, y, a, b = net["x"], net["y"], net["a"][edge_idx], net["b"][edge_idx]
    lines = shapely.linestrings(np.stack([np.c_[x[a], y[a]], np.c_[x[b], y[b]]], axis=1))
    return shapely.simplify(shapely.union_all(shapely.buffer(lines, 90)), 30)


def merge_by_name(areas):
    """One entry per area name: separate pockets with the same name are combined."""
    merged = {}
    for ar in areas:
        m = merged.get(ar["name"])
        if m is None:
            merged[ar["name"]] = {**ar, "pockets": 1}
            continue
        for k in ("residents", "flooded", "pre_cap", "pot_n", "pot_cap_est", "pot_no_est", "care_n", "care_residents",
                  "dialysis_n", "fire_n", "hospital_n"):
            m[k] += ar[k]
        m["pre_dry"] += ar["pre_dry"]
        m["pre_wet"] += ar["pre_wet"]
        m["cut_roads"] = list(dict.fromkeys(m["cut_roads"] + ar["cut_roads"]))[:6]
        m["outline"] = shapely.union_all([m["outline"], ar["outline"]])
        m["pockets"] += 1
    return sorted(merged.values(), key=lambda x: -x["residents"])


def cutoff_analysis(rel, transform):
    net = road_network()
    n, a, b = len(net["x"]), net["a"], net["b"]
    lab0 = components(n, a, b, np.ones(len(a), bool))
    main0 = np.bincount(lab0).argmax()
    in_main0 = lab0 == main0
    main_nodes = np.flatnonzero(in_main0)
    tree = cKDTree(np.c_[net["x"][main_nodes], net["y"][main_nodes]])

    z = np.load(WORK / "block_flood.npz")
    blocks = block_shapes(z["x"], z["y"])
    bd, bi = tree.query(np.c_[z["x"], z["y"]], distance_upper_bound=SNAP_M)
    bnode = np.where(np.isfinite(bd), main_nodes[np.minimum(bi, len(main_nodes) - 1)], -1)
    bpt = gpd.GeoSeries(gpd.points_from_xy(z["x"], z["y"]), crs=CRS).to_crs(4326)
    blon, blat = bpt.x.values, bpt.y.values
    pop = blocks["pop"].values.astype(float)
    # flooded share of each block's land, per variant and level (blocks outside the box never flood)
    block_id = rasterio.features.rasterize(((g, i + 1) for i, g in enumerate(blocks.to_crs(4326).geometry)),
                                           out_shape=rel.shape, transform=transform, fill=0, dtype="int32")
    land = (rel > LAND_FT) & (block_id > 0)
    land_cells = np.bincount(block_id[land], minlength=len(blocks) + 1)[1:]

    feats = json.load(open(DOCS / "points.json"))["features"]
    cap = json.load(open(DOCS / "capacity.json"))
    plon = np.array([f["geometry"]["coordinates"][0] for f in feats])
    plat = np.array([f["geometry"]["coordinates"][1] for f in feats])
    ppt = gpd.GeoSeries(gpd.points_from_xy(plon, plat), crs=4326).to_crs(CRS)
    pd_, pi_ = tree.query(np.c_[ppt.x.values, ppt.y.values], distance_upper_bound=SNAP_M)
    pnode = np.where(np.isfinite(pd_), main_nodes[np.minimum(pi_, len(main_nodes) - 1)], -1)
    pcols, prows = ~transform * (plon, plat)
    prows, pcols = np.floor(prows).astype(int), np.floor(pcols).astype(int)
    in_grid = (prows >= 0) & (prows < rel.shape[0]) & (pcols >= 0) & (pcols < rel.shape[1])

    result = {"variants": {}, "site_reach": {}}
    road_levels = {}
    for variant, fname in VARIANTS.items():
        reach = np.load(OUT / fname)
        lvl = edge_cut_levels(net, rel, transform, reach)
        road_levels[variant] = lvl
        site_reach = np.where(in_grid, reach[np.clip(prows, 0, rel.shape[0] - 1), np.clip(pcols, 0, rel.shape[1] - 1)], np.inf)
        for f, r in zip(feats, site_reach):
            if np.isfinite(r) and r < 10:
                result["site_reach"].setdefault(str(f["properties"]["id"]), {})[variant] = round(float(r), 2)
        steps = {}
        for L in LEVELS:
            flooded_cells = np.bincount(block_id[land & (reach <= L)], minlength=len(blocks) + 1)[1:]
            share = np.divide(flooded_cells, land_cells, out=np.zeros(len(blocks)), where=land_cells > 0)
            flooded = pop * share
            cut = lvl <= L
            lab = components(n, a, b, ~cut)
            mainL = np.bincount(lab[main_nodes]).argmax()
            stranded_node = in_main0 & (lab != mainL)
            # attach homes and sites to the nearest road point that still has a dry road at this level
            dry_node = np.zeros(n, bool)
            dry_node[a[~cut]] = dry_node[b[~cut]] = True
            usable = main_nodes[dry_node[main_nodes]]
            t2 = cKDTree(np.c_[net["x"][usable], net["y"][usable]])
            d1, i1 = t2.query(np.c_[z["x"], z["y"]], distance_upper_bound=SNAP_M)
            bnode = np.where(np.isfinite(d1), usable[np.minimum(i1, len(usable) - 1)], -1)
            d2, i2 = t2.query(np.c_[ppt.x.values, ppt.y.values], distance_upper_bound=SNAP_M)
            pnode = np.where(np.isfinite(d2), usable[np.minimum(i2, len(usable) - 1)], -1)
            b_cut = (bnode >= 0) & stranded_node[np.maximum(bnode, 0)]
            p_cut = (pnode >= 0) & stranded_node[np.maximum(pnode, 0)]
            areas = []
            for comp in np.unique(lab[stranded_node]):
                nodes = stranded_node & (lab == comp)
                in_b = b_cut & (lab[np.maximum(bnode, 0)] == comp)
                residents = pop[in_b].sum()
                if residents < MIN_AREA_RESIDENTS:
                    continue
                in_p = np.flatnonzero(p_cut & (lab[np.maximum(pnode, 0)] == comp))
                site_dry = lambda i: site_reach[i] > L
                pre_dry = [int(feats[i]["properties"]["id"]) for i in in_p if feats[i]["properties"]["k"] == "pre" and site_dry(i)]
                pre_wet = [int(feats[i]["properties"]["id"]) for i in in_p if feats[i]["properties"]["k"] == "pre" and not site_dry(i)]
                pot = [i for i in in_p if feats[i]["properties"]["k"] in POTENTIAL and site_dry(i)]
                pot_cap = [cap["types"].get(feats[i]["properties"].get("b"), {}).get("overnight") for i in pot]
                kinds = lambda k: [i for i in in_p if feats[i]["properties"]["k"] == k]
                edge_in = nodes[a] | nodes[b]
                boundary = np.flatnonzero((nodes[a] ^ nodes[b]) & cut)
                road_names = [nm for nm in net["name"][boundary] if nm]
                areas.append({
                    "name": subarea(blon[in_b], blat[in_b], pop[in_b]),
                    "residents": int(round(residents)), "flooded": int(round(flooded[in_b].sum())),
                    "pre_dry": pre_dry, "pre_wet": pre_wet,
                    "pre_cap": int(sum(cap["sites"].get(str(i), [0])[0] for i in pre_dry)),
                    "pot_n": len(pot), "pot_cap_est": int(sum(c for c in pot_cap if c)), "pot_no_est": sum(1 for c in pot_cap if not c),
                    "care_n": len(kinds("care")), "care_residents": int(sum(feats[i]["properties"].get("cap", 0) for i in kinds("care"))),
                    "dialysis_n": len(kinds("dialysis")), "fire_n": len(kinds("fire")), "hospital_n": len(kinds("hospital")),
                    "cut_roads": [nm for nm, _ in sorted(((nm, road_names.count(nm)) for nm in set(road_names)), key=lambda x: -x[1])][:6],
                    "outline": edges_outline(net, np.flatnonzero(edge_in & ~cut)),
                })
            areas = merge_by_name(areas)
            in_focus_edges = np.isfinite(lvl)
            cut_km = float((np.hypot(net["x"][a] - net["x"][b], net["y"][a] - net["y"][b])[cut & in_focus_edges]).sum() / 1000)
            steps[f"{L:.1f}"] = {"flooded": int(round(flooded.sum())), "flooded_main": int(round(flooded[~b_cut].sum())),
                                 "roads_cut_km": round(cut_km, 1), "areas": areas}
            print(f"    {variant} +{L:.1f} ft: flooded {flooded.sum():,.0f}, roads cut {cut_km:.1f} km, cut-off areas: "
                  + "; ".join(f"{x['name']} {x['residents']:,} res ({x['flooded']:,} flooded), shelter inside {x['pre_cap']}" for x in areas), flush=True)
        result["variants"][variant] = steps
    return result, net, road_levels


# ---- publish ---------------------------------------------------------------------------------

def write_png(mask, path, rgba):
    from PIL import Image
    img = np.zeros(mask.shape + (4,), "uint8")
    img[mask] = rgba
    Image.fromarray(img, "RGBA").save(path, optimize=True)


def publish(rel, transform, result, net, road_levels):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    WEB.mkdir(parents=True, exist_ok=True)
    for sub_dir in ("flood", "art"):
        (WEB / sub_dir).mkdir(exist_ok=True)
    step = 3  # images at ~9 m
    rel_s = rel[::step, ::step]
    wet = rel_s > WATER_FT
    for variant, fname in VARIANTS.items():
        reach = np.load(OUT / fname)[::step, ::step]
        for L in LEVELS:
            write_png(wet & (reach <= L), WEB / "flood" / f"{variant}_{L:.1f}.png", (42, 120, 214, 175))
    art = np.load(OUT / "art_grids.npz")
    for inch in ART_LAYERS:
        write_png(art[f"in{inch}"][::step, ::step] & wet, WEB / "art" / f"art_{inch}in.png", (232, 104, 52, 150))
    h, w = rel.shape
    corners = [transform * (0, 0), transform * (w, 0), transform * (w, h), transform * (0, h)]

    # impassable road segments (either variant, up to +3.0 ft), merged by road name and cut levels
    lv_m, lv_h = road_levels["mapped"], road_levels["hold"]
    show = np.flatnonzero(np.minimum(lv_m, lv_h) <= LEVELS[-1] + 1e-6)
    groups = {}
    for i in show:
        key = (net["name"][i], round(float(lv_m[i]), 1) if lv_m[i] <= 3.0 else None, round(float(lv_h[i]), 1) if lv_h[i] <= 3.0 else None)
        groups.setdefault(key, []).append(i)
    rows = []
    for (name, cm, ch), idx in groups.items():
        idx = np.array(idx)
        a, b = net["a"][idx], net["b"][idx]
        lines = shapely.linestrings(np.stack([np.c_[net["lon"][a], net["lat"][a]], np.c_[net["lon"][b], net["lat"][b]]], axis=1))
        rows.append({"n": name, "cm": cm, "ch": ch, "geometry": shapely.line_merge(shapely.union_all(lines))})
    roads = gpd.GeoDataFrame(rows, geometry="geometry", crs=4326)
    (WEB / "roads.geojson").unlink(missing_ok=True)
    roads.to_file(WEB / "roads.geojson", driver="GeoJSON", COORDINATE_PRECISION=5)

    # cut-off area outlines
    out_rows, steps = [], {}
    for variant, by_level in result["variants"].items():
        steps[variant] = {}
        for L, st in by_level.items():
            areas = []
            for ar in st["areas"]:
                props = {k: v for k, v in ar.items() if k != "outline"}
                areas.append(props)
                out_rows.append({"variant": variant, "level": float(L), "name": ar["name"], "residents": ar["residents"],
                                 "flooded": ar["flooded"], "pre_cap": ar["pre_cap"], "geometry": ar["outline"]})
            steps[variant][L] = {**{k: v for k, v in st.items() if k != "areas"}, "areas": areas}
    outlines = gpd.GeoDataFrame(out_rows, geometry="geometry", crs=CRS).to_crs(4326) if out_rows else None
    (WEB / "areas.geojson").unlink(missing_ok=True)
    if outlines is not None:
        outlines.to_file(WEB / "areas.geojson", driver="GeoJSON", COORDINATE_PRECISION=5)

    barriers = json.loads((OUT / "barriers.json").read_text())
    spill_feats = [{"type": "Feature", "geometry": {"type": "Point", "coordinates": [s_["lon"], s_["lat"]]},
                    "properties": {"crest": s_["crest_ft"], "km2": s_["land_km2"], "held": int(any(abs(w_["lat"] - s_["lat"]) < 1e-4 and abs(w_["lon"] - s_["lon"]) < 1e-4 for w_ in barriers["walls"]))}}
                   for s_ in barriers["spills"]]
    walls = [{"type": "Feature", "geometry": {"type": "Point", "coordinates": [w_["lon"], w_["lat"]]}, "properties": {"crest": w_["crest_ft"], "km2": w_["land_km2"], "held": 1}}
             for w_ in barriers["walls"]]
    (WEB / "spills.geojson").write_text(json.dumps({"type": "FeatureCollection", "features": spill_feats + walls}))
    meta = {"levels": LEVELS, "image_corners": [[round(x, 6), round(y, 6)] for x, y in corners], "focus": FOCUS,
            "depth_limit_ft": DEPTH_LIMIT_FT, "steps": steps, "site_reach": result["site_reach"],
            "art": json.loads((OUT / "art.json").read_text()), "calibration": json.loads((OUT / "calibration.json").read_text())}
    (WEB / "steps.json").write_text(json.dumps(meta, separators=(",", ":")))

    # private spreadsheet (contains capacities)
    wb = Workbook()
    ws = wb.active
    ws.title = "Steps"
    head = ["Variant", "Bay level (ft)", "Flooded residents (focus area)", "Flooded, not cut off", "Major + local roads impassable (km)",
            "Cut-off area", "Residents in area", "Flooded residents in area", "Pre-identified shelter capacity inside (dry)",
            "Potential sites inside (dry)", "Est. capacity of potential sites", "Care facilities inside", "Licensed care residents inside",
            "Dialysis clinics inside", "Fire stations inside", "Roads cutting it off"]
    ws.append(head)
    for variant, by_level in steps.items():
        vname = "As mapped" if variant == "mapped" else "If embankment holds"
        for L, st in by_level.items():
            base = [vname, float(L), st["flooded"], st["flooded_main"], st["roads_cut_km"]]
            if not st["areas"]:
                ws.append(base + ["(none)"])
            for ar in st["areas"]:
                ws.append(base + [ar["name"], ar["residents"], ar["flooded"], ar["pre_cap"], ar["pot_n"], ar["pot_cap_est"], ar["care_n"],
                                  ar["care_residents"], ar["dialysis_n"], ar["fire_n"], "; ".join(ar["cut_roads"])])
    sp = wb.create_sheet("Spill points")
    sp.append(["Opens at (ft)", "Crest (ft above MHHW)", "Land opened (km2)", "Lat", "Lon", "Blocked in 'holds' variant", "Map"])
    for s_ in spill_feats:
        lon, lat = s_["geometry"]["coordinates"]
        sp.append([None, s_["properties"]["crest"], s_["properties"]["km2"], lat, lon, "Yes" if s_["properties"]["held"] else "No",
                   f'=HYPERLINK("https://www.google.com/maps/@{lat},{lon},19z/data=!3m1!1e3","Satellite")'])
    for i, s_ in enumerate(barriers["spills"], 2):
        sp.cell(i, 1, s_["level"])
    ar_ws = wb.create_sheet("ART vs NOAA")
    ar_ws.append(["ART level (in)", "Variant", "ART flooded land (km2)", "NOAA-based flooded land (km2)", "Both (km2)", "Overlap (IoU)"])
    for r in meta["art"]:
        ar_ws.append([r["inches"], r["variant"], r["art_km2"], r["noaa_km2"], r["both_km2"], r["iou"]])
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=c.row == 1)
                if c.row == 1:
                    c.fill = PatternFill("solid", fgColor="DDEBF7")
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 18
    wb.save(OUT / "airport_thresholds_2_to_3ft.xlsx")


def main():
    t0 = time.time()
    OUT.mkdir(parents=True, exist_ok=True)
    start = sys.argv[1] if len(sys.argv) > 1 else "grid"
    stages = ["grid", "reach", "barriers", "art", "roads", "publish"]
    run = stages[stages.index(start):]

    if "grid" in run or not (OUT / "grid.npz").exists():
        rel, transform = build_grid()
        log(t0, f"grid {rel.shape}, {cell_km2(transform) * rel.size:.0f} km2")
    else:
        rel, transform = load_grid()
    bay = open_water(rel)

    if "reach" in run or not (OUT / "reach.npy").exists():
        cal = calibrate(rel, transform, bay)
        log(t0, "calibration against NOAA polygons (land only):")
        for name, rows in cal.items():
            print(f"    {name}: " + "  ".join(f"+{ft} ft ours {r['ours_km2']} / NOAA {r['noaa_km2']} km2, IoU {r['iou']}" for ft, r in rows.items()))
        best = max(cal, key=lambda n: np.mean([r["iou"] for r in cal[n].values()]))
        reach = water_reach(rel, bay, eight=best == "8-neighbor")
        np.save(OUT / "reach.npy", reach)
        (OUT / "calibration.json").write_text(json.dumps({"best": best, **cal}, indent=1))
        log(t0, f"water reach saved ({best})")
    reach = np.load(OUT / "reach.npy")

    if "barriers" in run or not (OUT / "reach_hold.npy").exists():
        spills = find_spills(reach, rel, transform, 0.03)
        reach_hold, walls = barriers_hold(rel, transform, bay)
        np.save(OUT / "reach_hold.npy", reach_hold)
        (OUT / "barriers.json").write_text(json.dumps({"spills": spills, "walls": walls}, indent=1))
        log(t0, f"{len(spills)} spill points; {len(walls)} blocked for the barriers-hold variant")
    reach_hold = np.load(OUT / "reach_hold.npy")

    if "art" in run or not (OUT / "art.json").exists():
        fetch_art()
        rows = compare_art(rel, transform, reach, reach_hold)
        (OUT / "art.json").write_text(json.dumps(rows, indent=1))
        log(t0, "ART comparison (land only):")
        for r in rows:
            print(f"    {r['inches']} in ({r['variant']}): ART {r['art_km2']} / NOAA {r['noaa_km2']} km2, both {r['both_km2']}, IoU {r['iou']}")

    if "roads" in run or not (OUT / "cutoffs.pkl").exists():
        import pickle
        result, net, road_levels = cutoff_analysis(rel, transform)
        with open(OUT / "cutoffs.pkl", "wb") as fh:
            pickle.dump({"result": result, "road_levels": road_levels}, fh)
        log(t0, "road cut-off analysis saved")
    import pickle
    with open(OUT / "cutoffs.pkl", "rb") as fh:
        saved = pickle.load(fh)
    publish(rel, transform, saved["result"], road_network(), saved["road_levels"])
    log(t0, f"published to {WEB.relative_to(DOCS.parent.parent)}")


if __name__ == "__main__":
    main()
